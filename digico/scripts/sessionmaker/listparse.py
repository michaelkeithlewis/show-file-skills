#!/usr/bin/env python3
"""listparse.py - read an input list (any common format) into rows and channels.

The deterministic half of intake: no network, no model calls.

Sources it reads (load_rows_from(filename, data, text)):
  .csv / .tsv / .txt-with-tabs   delimited grids, header row found automatically
  .xlsx / .xlsm                  every sheet scored, the best input table wins (openpyxl when it
                                 imports, else a built-in zip/XML reader)
  .xls                           Excel 97-2003 (xlrd when it imports, else the built-in BIFF8 reader
                                 below); an .xls that is really HTML or text is read as such
  .docx                          Word tables (best input table wins), else the paragraphs as text
  .pdf                           text layer through `pdftotext -layout` (poppler); one or two columns
  .htm / .html                   <table> rows
  pasted text                    "1 Kick In Beta91 48V", "- Snare top (SM57)", "8-9 OH L/R C414 +48"

Pipeline: rows -> expand_pairs() ("OH L/R" -> two rows) -> pair_rows() (L/R neighbours -> one stereo
channel) -> build_channels() (8-char names, 48V, source, coarse category, confidence, flags).
Categories here are the parser's 28 coarse ids; sessionmaker.taxonomy refines them to the 50
template ids.

As a module: load_rows_from(...) / load_rows(path) -> (rows, meta); parse_file / parse_text -> request.
CLI: python3 -m sessionmaker.listparse LIST [--table]
"""
import csv
import io
import json
import os
import re
import struct
import subprocess
import sys
import unicodedata
import zipfile
from html.parser import HTMLParser
from xml.etree import ElementTree as ET

SCHEMA_VERSION = 1
MAX_NAME = 8


HEADER_FIELDS = [
    ('short', r'^(name ?\(8\)|short ?name|console ?name|scribble|8 ?char.*|desk ?name|ses ?name)$'),
    ('ch', r'^(ch|ch\.|chan|channel|ch ?#|ch ?no\.?|#|no\.?|in ?#|input ?#|inp ?#|foh ?ch|mon ?ch|line|split|number|num|n)$'),
    ('p48', r'^(\+?48 ?v?|phantom( power)?|p48|ph|pp|48v ?\?|\+48v)$'),
    ('stereo', r'^(st|st\.|st\?|stereo|stereo\?|mono ?/ ?st(ereo)?|st(ereo)? ?/ ?mono|m ?/ ?s|l ?/ ?r|lr|pair|paired)$'),
    ('mic', r'^(mic|mics|mic ?/ ?di|mic ?/ ?di ?/ ?line|mic ?di|microphone|di|di ?/ ?mic|mic ?/ ?line|transducer|pickup|'
            r'device|mic ?or ?di|mic/di/source|.*show mic.*|mic ?model|mic ?type)$'),
    ('stand', r'^(stand|stands|mount|stand ?type|clip ?/ ?stand)$'),
    ('notes', r'^(note|notes|comment|comments|remarks|info|description|desc)$'),
    ('who', r'^(who|player|musician|performer|owner|band ?member)$'),
    ('name', r'^(input|inputs|input ?name|name|source|instrument|inst|channel ?name|master ?input|signal|item|'
             r'input ?/ ?source|what|line ?name|label)$'),
    ('patch', r'^(sub|sub ?snake|subsnake|stage ?box|stagebox|box|patch|socket|fest\.? ?patch|snake|mon( dm.*)?|'
              r'foh ?rpio|rack|port|nc14|dm-?\d+.*|.*rio.*|sd-?rack.*|stage ?rack.*|sub ?slot|fan)$'),
    ('xlr', r'^(xlr|cable|length|xlr ?length|xlr ?len)$'),
    ('gain', r'^(gain|gain ?\(db\)|trim)$'),
    ('pos', r'^(pos\.?|position|location|loc|where)$'),
]
HEADER_RE = [(f, re.compile(rx)) for f, rx in HEADER_FIELDS]
FIELD_RX = dict(HEADER_FIELDS)
OUTPUT_HEADER = re.compile(r'^(out|outs|output|outputs|output ?name|mix|iem|aux)$')
STOP_NAME = re.compile(r'^(counts?|item|outputs?|totals?|loops|notes?:?|patch notes|key|legend|output list)$', re.I)

MIC_BRANDS = [
    'shure', 'sennheiser', 'senn', 'beyer', 'beyerdynamic', 'beyerdnamic', 'audix', 'akg', 'neumann', 'dpa', 'rode',
    'audio-technica', 'audio technica', 'electro-voice', 'electro voice', 'telefunken', 'royer', 'coles', 'mojave',
    'schoeps', 'earthworks', 'josephson', 'sony', 'heil', 'lewitt', 'warm audio', 'aston', 'radial', 'countryman', 'bss',
    'klark', 'palmer', 'avalon', 'rupert neve', 'rndi', 'sansamp', 'livewire', 'whirlwind', 'ev', 'at', 'sennheisser',
]
MIC_MODELS = [
    r'beta ?\d{2,3}a?', r'sm ?\d{2}(?:a|b|s)?', r'ksm ?\d{2,3}', r'pga ?\d{2}', r'e ?\d{3}', r'md ?\d{3}', r'mkh ?\d{3}',
    r'mke ?\d', r'km ?\d{3}', r'kms ?\d{3}', r'tlm ?\d{2,3}', r'u ?87', r'u ?47', r'c ?\d{3}b?', r'd ?112', r'd\d{1,2}',
    r'i5', r'm ?\d{2,3}(?: ?tg)?', r'tg ?[dm]\d{2,3}c?', r're ?\d{2,3}', r'nd ?\d{2,3}', r'pr ?\d{2}', r'at ?\d{4}',
    r'r-?1[02]1', r'4038', r'vp ?88', r'v ?7', r'om ?\d', r'b\d{2}a?', r'\d{4}v?', r'cm ?\d', r'sr ?25',
    r'57', r'58', r'91a?', r'52a?', r'98', r'421', r'414', r'604', r'904', r'906', r'609', r'184', r'81',
    r'ax(?:ient)?', r'ad ?[12]', r'ulx-?d', r'qlx-?d', r'ur ?4d', r'ew ?\d{3}',
]
MIC_GENERIC = [
    r'd\.?i\.?', r'di box', r'j ?48', r'jdi', r'pro ?48', r'pro ?d2', r'ldi', r'usb ?di', r'stereo ?di', r'active ?di',
    r'passive ?di', r'dual ?di', r'line', r'xlr', r'direct', r'rf', r'wireless', r'reamp(?: box)?', r'trs', r'madi', r'dante',
    r'aes', r'optogate', r'shotgun', r'lav(?:alier)?', r'headset',
]
GENERIC_RX = re.compile(r'(?i)^(?:' + '|'.join(MIC_GENERIC) + r')$')
MIC_RX = re.compile(r'(?i)(?<![\w.])(?:' + '|'.join(re.escape(b) for b in MIC_BRANDS) + '|' + '|'.join(MIC_MODELS)
                    + '|' + '|'.join(MIC_GENERIC) + r')(?![\w])')
P48_RX = re.compile(r'(?i)(?<![\w.])(?:\+?48 ?v(?:olts?)?|\+48|p48|phantom(?: power)?|48v on)(?![\w])')
NO48_RX = re.compile(r'(?i)\bno ?(?:\+?48 ?v?|phantom)\b|\bphantom ?off\b')
STAND_RX = re.compile(r'(?i)(?<![\w])(?:(?:tall|short|small|medium|med|heavy|mini|low|desk|kick|round[- ]base|weighted)\s?'
                      r'(?:boom|straight)(?:\s?stand)?(?:\s?\(heavy\))?|boom(?:\s?stand)?|straight(?:\s?stand)?|lp ?claw|claw|clip|'
                      r'clamp|magic ?arm|in ?drum|kelly ?shu|drum ?mount|piano ?mount|gooseneck|tripod|desk ?stand|'
                      r'hang(?:ing)?|taped|on ?amp|over ?cab|stereo ?bar|round ?base)(?![\w])')
SIDE_TOKENS = {'l': 'L', 'left': 'L', 'lft': 'L', 'lt': 'L', '(l)': 'L', 'r': 'R', 'right': 'R', 'rt': 'R', 'rgt': 'R', '(r)': 'R'}
STAGE_SIDE = {'sl': 'L', 'sr': 'R', 'stage left': 'L', 'stage right': 'R'}
PAIR_WORDS = re.compile(r'(?i)\s*(?:\(?\b(?:l ?[/&+] ?r|left ?[/&+] ?right|stereo|st\.?|pair)\b\)?)\s*')
STEREO_CELL = re.compile(r'(?i)^\s*(?:y|yes|x|✓|✔|st|st\.|stereo|2|l ?/ ?r|lr|true|1|pair|paired)\s*$')
STEREO_NOTE = re.compile(r'(?i)\bstereo\b|\bst pair\b|\bl ?/ ?r\b')
MIC_PAIR = re.compile(r'(?i)(?:\bx ?2\b|\b2 ?x\b|\bpair of\b|\bmatched pair\b|\bpair\b)')
STEREO_NATURE = ('overhead', 'audience', 'drum_room', 'keys', 'playback', 'music_playback', 'choir', 'drum_pad')

CATEGORIES = [
    'kick', 'snare', 'tom', 'hihat', 'overhead', 'drum_room', 'drum_trigger', 'drum_pad', 'percussion',
    'bass', 'electric_guitar', 'acoustic_guitar', 'keys', 'strings', 'brass_wind',
    'lead_vocal', 'backing_vocal', 'choir', 'speech', 'talkback',
    'playback', 'click_timecode', 'music_playback', 'audience', 'fx_return', 'utility', 'spare', 'other',
]
W = r'(?<![a-z0-9])'
E = r'(?![a-z])'
CAT_RULES = [
    ('spare', rf'^\s*(?:\(?spare\)?|spr|-+|—|hole|n/?a|empty|unused|tbd|ch ?\d+|input ?\d+)\s*$|{W}spares{E}', 0.95),
    ('talkback', rf'{W}(?:tb|talk ?back|talkback|t2me|t2\w*|comms?|shout|cam talk|foh ?> ?band|prod tb)(?![a-z])', 0.95),
    ('utility', rf'{W}(?:pink(?: noise)?|noise|ref(?:erence)? ?mic|ref|smaart|rta|measurement|meas|sig ?gen|osc(?:illator)?|'
                rf'tone|tuner|tune rig out|tech tune|gtr tune)(?![a-z])', 0.9),
    ('fx_return', rf'{W}(?:verb|reverb|rvb|delay|dly|slap|adt|chorus|thicken\w*|plate|hall|fx ?rt|fx ?ret\w*|fx ?rtn|'
                  rf'\w+ ?rt$|rtn|return|throw|lush|doubler|echo|ear ?verb|iem ?verb|gated v|drum v|vox v)(?![a-z])', 0.85),
    ('click_timecode', rf'{W}(?:click|clk|guide|cues?|count ?in|smpte|ltc|time ?code|mtc|midi|slate)(?![a-z])', 0.9),
    ('playback', rf'(?:^t\.|^t |{W}(?:trax|trx|tracks?|pb|playback|stems?|ableton|backing ?tracks?|seq|vsc|mtx))(?![a-z])', 0.85),
    ('music_playback', rf'{W}(?:ipod|walk ?in|walk ?out|walkin|music|pre ?show|gfx|video|vid|laptop|computer|mac|radio|'
                       rf'media|spotify|playout|2 ?mix|2mx|break ?music)(?![a-z])', 0.8),
    ('drum_trigger', rf'{W}trig(?:ger)?s?{E}', 0.95),
    ('drum_pad', rf'{W}(?:spd\w*|sample ?pad|pads?|sp-?404|octapad|sampler|mpc|drum ?pad|e-?kit|spdx)(?![a-z])', 0.85),
    ('choir', rf'{W}choir{E}', 0.95),
    ('speech', rf'{W}(?:mc|host|pastor|preacher|podium|lectern|announcer|speaker|sprkr|speak|spkr|script|panel|walkabout|'
               rf'lav|presenter|announce\w*|annc|vog|sermon|teaching|welcome)(?![a-z])', 0.8),
    ('lead_vocal', rf'{W}(?:lv|lead ?vo(?:x|cal)|main ?vo(?:x|cal)|ld ?vox|artist ?vo(?:x|cal))(?![a-z])', 0.9),
    ('backing_vocal', rf'{W}(?:bgv|bgvs|bv|bvs|harmony|harm|backing ?vo\w*|background ?vo\w*|bkg ?vo\w*)(?![a-z])', 0.9),
    ('vocal', rf'{W}(?:vox|vocals?|voc|vx|singer|sing|headset|vocal mic)(?![a-z])', 0.65),
    ('kick', rf'{W}(?:kick|kik|kck|bd|bass ?drum|sub ?kick|k ?in|k ?out){E}', 0.95),
    ('snare', rf'{W}(?:snare|snr|sn|sd|rim){E}', 0.9),
    ('hihat', rf'{W}(?:hi-?hats?|hh|hats?|hihats?){E}', 0.95),
    ('tom', rf'{W}(?:toms?|rack|floor|ft|rt|roto|big ?floor){E}', 0.85),
    ('overhead', rf'{W}(?:oh|ohs|overheads?|cymbals?|cym|ride|crash|china|splash|crotch){E}', 0.9),
    ('drum_room', rf'{W}(?:drum ?room|room|vibe|kit ?mic|front ?of ?kit|trash|crush|drum ?mic|d\. ?room|d\. ?vibe|drum ?amb\w*){E}', 0.75),
    ('percussion', rf'{W}(?:perc\w*|conga\w*|bongos?|cajon|djembe|shaker|tamb\w*|timbales?|cowbell|clap|glock\w*|vibraphone|'
                   rf'marimba|xylo\w*|chimes|gong|timp\w*|washboard|cabasa|guiro|wood ?block|cajón|darbuka|tabla){E}', 0.9),
    ('brass_wind', rf'{W}(?:trumpet|tpt|trp|sax\w*|alto|tenor|bari|trombone|tbn|bone|horns?|tuba|flute|clarinet|oboe|bassoon|'
                   rf'harmonica|brass|winds?|woodwinds?|flugel\w*|cornet|whistle|ewi){E}', 0.9),
    ('strings', rf'{W}(?:violins?|vln|viola|vla|cello|vc|fiddle|strings?|str\d|harp|quartet){E}', 0.85),
    ('acoustic_guitar', rf'{W}(?:ag|acoustic|acu|acous|ac ?gtr|nylon|classical|12 ?string|mando\w*|banjo|ukulele|uke|bouzouki|'
                        rf'resonator|dobro){E}', 0.85),
    ('bass', rf'{W}(?:bass|bs|b\. ?di|b\. ?mic|bass ?di|sub ?bass|synth ?bass|bass ?synth|upright|double ?bass|keybass|bgtr){E}', 0.9),
    ('electric_guitar', rf'{W}(?:eg|egt|eg ?\d|elec\w*|ele|gtr|guitars?|amp|cab|kemper|helix|axe ?fx|fractal|quad ?cortex|qc|'
                        rf'strat|tele|marshall|lap ?steel|pedal ?steel|steel|slide|baritone|bari ?gtr){E}', 0.8),
    ('keys', rf'{W}(?:keys?|kb|kbd|keyboards?|piano|pno|nord|rhodes|wurli\w*|organ|b3|hammond|leslie|synths?|syn|prophet|juno|'
             rf'moog|korg|dx7|cp\d+|mellotron|mello|clav|op-?1|keylargo|grand|mainstage|accordion|harmonium|ob-?6|ob-?x|'
             rf'rhodes|minilogue|kemper keys|ivory|vintage keys|pad synth){E}', 0.85),
    ('audience', rf'{W}(?:aud|audience|crowd|crowds|amb|ambient|ambience|house ?mic|shotgun){E}', 0.9),
]
CAT_RULES = [(c, re.compile(rx), p) for c, rx, p in CAT_RULES]
SECTION_PRIOR = [
    (re.compile(r'(?i)drum|kit|riser'), 'drum_room'), (re.compile(r'(?i)track|trax|playback|pb'), 'playback'),
    (re.compile(r'(?i)spd|pad'), 'drum_pad'), (re.compile(r'(?i)vocal|vox|bgv'), 'backing_vocal'),
    (re.compile(r'(?i)keys|keyboard'), 'keys'), (re.compile(r'(?i)perc'), 'percussion'),
    (re.compile(r'(?i)choir'), 'choir'), (re.compile(r'(?i)horn|brass'), 'brass_wind'), (re.compile(r'(?i)string'), 'strings'),
    (re.compile(r'(?i)fx|effect|return'), 'fx_return'), (re.compile(r'(?i)aud|crowd|amb'), 'audience'),
]
MIC_PRIOR = [
    (re.compile(r'(?i)\bbeta ?91|\bbeta ?52|\bd ?112\b|\bd6\b|\be ?90[12]\b|kelly ?shu|\btg ?d71|\bd71\b|\bre ?20\b|\b4055\b'), 'kick'),
    (re.compile(r'(?i)\be ?60[49]\b|\be ?904\b|\bbeta ?98|\bd4\b|\bd2\b|\bmd ?421\b'), 'tom'),
    (re.compile(r'(?i)\be ?906\b|\be ?609\b|\br-?121\b|\bmd ?409\b'), 'electric_guitar'),
    (re.compile(r'(?i)vp ?88'), 'drum_room'),
    (re.compile(r'(?i)mkh ?41[06]|at ?4053|at ?897|shotgun'), 'audience'),
    (re.compile(r'(?i)4018|4099|beta ?58|sm ?58|ksm ?[89]|v ?7\b|ksm ?11|om ?\d|2028|ksm9|e ?935|e ?945'), 'vocal'),
]

PHRASES = [
    (r'(?:acoustic|acous|ac) (?:guitar|gtr)', 'AG'), (r'(?:electric|elec) (?:guitar|gtr)', 'EG'), (r'bass guitar', 'Bass'), (r'lead vocals?', 'LV'),
    (r'(?:background|backing|bkg) vocals?', 'BGV'), (r'talk ?back', 'TB'), (r'hi-? ?hats?', 'Hat'),
    (r'floor tom', 'Floor'), (r'rack tom', 'Rack'), (r'kick drum', 'Kick'), (r'bass drum', 'Kick'), (r'drum pad', 'Pad'),
    (r'stage left', 'SL'), (r'stage right', 'SR'), (r'down ?stage', 'DS'), (r'up ?stage', 'US'), (r'pink noise', 'Pink'),
    (r'time ?code', 'LTC'), (r'walk[- ]?in', 'WalkIn'), (r'synth bass', 'SynBass'), (r'bass synth', 'BassSyn'),
    (r'drum room', 'D.Room'), (r'pedal steel', 'PdlSteel'), (r'lap steel', 'LapSteel'), (r'reference mic', 'Ref'),
    (r'measurement mic', 'Meas'), (r'sample pad', 'Pad'), (r'audience', 'Aud'), (r'overheads?', 'OH'),
]
PHRASES = [(re.compile(r'(?i)\b' + p + r'\b'), r) for p, r in PHRASES]
WORDS = {
    'snare': 'Snr', 'bottom': 'Bot', 'btm': 'Bot', 'bttm': 'Bot', 'vocal': 'Vox', 'vocals': 'Vox', 'acoustic': 'Acu',
    'electric': 'Ele', 'guitar': 'Gtr', 'guitars': 'Gtrs', 'keyboard': 'Keys', 'keyboards': 'Keys', 'talkback': 'TB',
    'trigger': 'Trig', 'triggers': 'Trig', 'tracks': 'Trax', 'track': 'Trax', 'playback': 'PB', 'piano': 'Pno', 'organ': 'Org',
    'synth': 'Syn', 'percussion': 'Perc', 'spare': 'Spr', 'guest': 'Gst', 'trumpet': 'Tpt', 'trombone': 'Tbn',
    'saxophone': 'Sax', 'violin': 'Vln', 'mandolin': 'Mando', 'wireless': 'RF', 'microphone': 'Mic', 'left': 'L',
    'right': 'R', 'center': 'C', 'centre': 'C', 'condenser': 'Cond', 'dynamic': 'Dyn', 'ribbon': 'Rib', 'amplifier': 'Amp',
    'monitor': 'Mon', 'crowd': 'Crowd', 'ambient': 'Amb', 'ambience': 'Amb', 'background': 'BG', 'harmony': 'Harm',
    'overhead': 'OH', 'cymbal': 'Cym', 'cymbals': 'Cym', 'shaker': 'Shkr', 'tambourine': 'Tamb', 'congas': 'Cnga',
    'bongos': 'Bngo', 'djembe': 'Djmb', 'cajon': 'Cajon', 'click': 'Click', 'smpte': 'SMPTE', 'master': 'Mst',
    'sub': 'Sub', 'main': 'Main', 'reverb': 'Verb', 'delay': 'Dly', 'return': 'Rtn', 'audience': 'Aud', 'mellotron': 'Mello',
    'prophet': 'Proph', 'hammond': 'B3', 'leslie': 'Leslie', 'wurlitzer': 'Wurli', 'rhodes': 'Rhodes', 'timpani': 'Timp',
    'speaker': 'Spkr', 'announcer': 'Annc', 'pastor': 'Pastor', 'podium': 'Podium', 'lectern': 'Lectrn', 'presenter': 'Prsntr',
    'computer': 'Comp', 'laptop': 'Laptop', 'video': 'Video', 'music': 'Music', 'stereo': 'St', 'mono': 'Mono',
    'upright': 'Upright', 'baritone': 'Bari', 'tenor': 'Tenor', 'alto': 'Alto', 'trombones': 'Tbns', 'horns': 'Horns',
    'strings': 'Strs', 'string': 'Str', 'cello': 'Cello', 'viola': 'Vla', 'fiddle': 'Fiddle', 'banjo': 'Banjo',
    'resonator': 'Reso', 'measurement': 'Meas', 'reference': 'Ref', 'option': '', 'optional': '', 'opt': '',
}
DROP_WORDS = re.compile(r'(?i)\b(?:option(?:al)?|tbc|tbd|maybe|potentially|possibly|if needed)\b[:.]?|\?')
NOT_PERSON = set(w.lower() for w in (
    'Kick Snare Snr Tom Toms Rack Floor Hat Hats Ride Crash OH Overhead Room Vibe Bass Sub Synth Keys Key Nord Piano Organ Pno '
    'Acoustic Electric Ele Gtr Guitar Amp Vox Vocal Lead Main Spare Spr Guest Gst Click Guide Cues Track Tracks Trax PB Playback '
    'Crowd Audience Aud Amb Pink Ref Talkback TB FOH MON Stage Drum Drums Pad Pads SPD Trig Trigger Perc Percussion Choir Horn '
    'Horns Trumpet Sax Trombone String Strings Violin Cello Viola Fiddle Mando Banjo Steel Lap Pedal Leslie Prophet Juno Moog '
    'Rhodes Wurli Mellotron Kemper Helix Reamp Dyn Rib Ribbon Dyn Top Bot Bottom In Out Low High Hi Lo Big Small Left Right '
    'Center Centre Up Down Wet Dry Clean Dirty DI Mic Line RF Option Opt Walk Music Video Radio Smpte LTC Timecode Shout '
    'Comms Pastor Host MC Speaker Podium Lectern Announcer Script Panel Walkabout Headset Lav Harmony BGV BV LV Verb Delay '
    'Slap Plate Hall Return Rtn FX Tune Tuner Trax Stereo Mono Kit Front Back Side Bari Tenor Alto Upright Grand Clap '
    'Shaker Tamb Conga Congas Bongo Bongos Cajon Djembe Timbale Glock Vibes Marimba Gong Ableton Laptop Computer IPod '
    'Audio Sample Sampler Ambient Ambience Fill Spot Wedge Mix Cam Camera Prod Tech Rig Tone Noise Measurement '
    'Dark Bright Mid Mids Wide Close Far Near Shell Inside Outside Hits Trx SL SR USC DSC USR USL DSR DSL C L R '
    'Talk Back Need Bring Laptop Stereo Mono Spare Main Lead Ensemble Worship Band Show Jazz'
    .split()))



def _cell_str(v):
    if v is None:
        return ''
    if hasattr(v, 'month') and hasattr(v, 'day'):
        return '%d/%d' % (v.month, v.day)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_xlsx(path):
    """[(sheet name, rows)] using openpyxl when present, else a minimal zip/XML reader.
    `path` may be a file path or a binary file object (io.BytesIO). Cells merged down over several rows are filled
    down (fill_merged_down)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [[_cell_str(v) for v in r] for r in ws.iter_rows(values_only=True)]
            out.append((ws.title, rows))
        wb.close()
    except ImportError:
        out = _read_xlsx_builtin(path)
    merges = _xlsx_merges(path)
    for name, rows in out:
        fill_merged_down(rows, merges.get(name, []))
    return out


XNS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
       'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
MERGE_RX = re.compile(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)"')


def _xlsx_sheets(z):
    """[(sheet name, part path)] in workbook order."""
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    target = {r.get('Id'): r.get('Target') for r in rels}
    out = []
    for sh in wb.find('m:sheets', XNS):
        p = target[sh.get('{%s}id' % XNS['r'])].lstrip('/')
        out.append((sh.get('name'), p if p.startswith('xl/') else 'xl/' + p))
    return out


def _col_num(letters):
    n = 0
    for ch in letters.decode() if isinstance(letters, bytes) else letters:
        n = n * 26 + ord(ch) - 64
    return n


def _xlsx_merges(src):
    """{sheet name: [(first row, first col, last row, last col)]} (1-based) from each sheet's <mergeCells>."""
    try:
        if hasattr(src, 'seek'):
            src.seek(0)
        z = zipfile.ZipFile(src)
        return {name: [(int(m[1]), _col_num(m[0]), int(m[3]), _col_num(m[2])) for m in MERGE_RX.findall(z.read(p))]
                for name, p in _xlsx_sheets(z)}
    except Exception:
        return {}


def fill_merged_down(rows, merges):
    """A cell merged down over several rows (a name over an L/R pair, a range "3-4" over its two rows) -> its value
    in each row of the merge, in the merge's first column only. Merges across one row (titles, section bands) are
    left alone, so a section heading stays a one-cell row."""
    for r0, c0, r1, _ in merges:
        if r1 <= r0 or r0 > len(rows) or c0 > len(rows[r0 - 1]) or not rows[r0 - 1][c0 - 1]:
            continue
        v = rows[r0 - 1][c0 - 1]
        for r in range(r0 + 1, min(r1, len(rows)) + 1):
            row = rows[r - 1]
            row.extend([''] * (c0 - len(row)))
            if not row[c0 - 1]:
                row[c0 - 1] = v


def _read_xlsx_builtin(path):
    ns = XNS
    if hasattr(path, 'seek'):
        path.seek(0)
    z = zipfile.ZipFile(path)
    shared = []
    if 'xl/sharedStrings.xml' in z.namelist():
        for si in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('m:si', ns):
            shared.append(''.join(t.text or '' for t in si.iter('{%s}t' % ns['m'])))
    out = []
    for sheet_name, p in _xlsx_sheets(z):
        rows = []
        for row in ET.fromstring(z.read(p)).iter('{%s}row' % ns['m']):
            cells = {}
            for c in row.findall('m:c', ns):
                ref = c.get('r')
                col = 0
                for ch in re.match(r'[A-Z]+', ref).group(0):
                    col = col * 26 + ord(ch) - 64
                v = c.find('m:v', ns)
                t = c.get('t')
                if t == 's' and v is not None:
                    val = shared[int(v.text)]
                elif t == 'inlineStr':
                    val = ''.join(x.text or '' for x in c.iter('{%s}t' % ns['m']))
                else:
                    val = v.text if v is not None else ''
                    try:
                        f = float(val)
                        val = str(int(f)) if f.is_integer() else val
                    except (TypeError, ValueError):
                        pass
                cells[col] = (val or '').strip()
            r = int(row.get('r'))
            while len(rows) < r - 1:
                rows.append([])
            rows.append([cells.get(i, '') for i in range(1, max(cells or {0: 0}) + 1)])
        out.append((sheet_name, rows))
    return out


def read_delimited(text):
    sample = text[:4000]
    if '\t' in sample:
        delim = '\t'
    else:
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=',;|').delimiter
        except csv.Error:
            delim = ','
    return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]


def pdf_text(path=None, data=None):
    """Text layer of a PDF, columns kept (`pdftotext -layout`). data= bytes go in on stdin."""
    if data is not None:
        r = subprocess.run(['pdftotext', '-layout', '-', '-'], input=data, capture_output=True, timeout=30)
        return r.stdout.decode('utf-8', 'replace')
    return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True, timeout=30).stdout


def decode_text(data):
    """bytes -> str: UTF-8 (with or without BOM), UTF-16 with BOM, else Windows-1252."""
    if data[:2] in (b'\xff\xfe', b'\xfe\xff'):
        return data.decode('utf-16', 'replace')
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        return data.decode('cp1252', 'replace')



OLE_MAGIC = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
_END = 0xFFFFFFFA


def _cfb_stream(data, want=('Workbook', 'Book')):
    """The bytes of the first stream named in `want` from an OLE2 compound file."""
    if data[:8] != OLE_MAGIC:
        raise ValueError('not an OLE2 compound file')
    ssz = 1 << struct.unpack_from('<H', data, 0x1E)[0]
    mssz = 1 << struct.unpack_from('<H', data, 0x20)[0]
    nfat, dir_start = struct.unpack_from('<II', data, 0x2C)
    cutoff, mfat_start, nmfat, difat_start, ndifat = struct.unpack_from('<IIIII', data, 0x38)
    per = ssz // 4

    def sec(i):
        return data[(i + 1) * ssz:(i + 2) * ssz]

    difat = list(struct.unpack_from('<109I', data, 0x4C))
    d = difat_start
    for _ in range(ndifat):
        if d >= _END:
            break
        v = struct.unpack('<%dI' % per, sec(d))
        difat.extend(v[:-1])
        d = v[-1]
    fat = []
    for i in difat[:nfat]:
        if i < _END:
            fat.extend(struct.unpack('<%dI' % per, sec(i)))

    def chain(start, table):
        out, k = [], start
        while k < _END and k < len(table) and len(out) <= len(table):
            out.append(k)
            k = table[k]
        return out

    dirs = b''.join(sec(i) for i in chain(dir_start, fat))
    entries = []
    for off in range(0, len(dirs) - 127, 128):
        nlen = struct.unpack_from('<H', dirs, off + 64)[0]
        name = dirs[off:off + max(0, nlen - 2)].decode('utf-16-le', 'replace')
        start, size = struct.unpack_from('<II', dirs, off + 116)
        entries.append((name, dirs[off + 66], start, size))
    root = entries[0]
    for name, typ, start, size in entries:
        if typ != 2 or name not in want:
            continue
        if size >= cutoff:
            return b''.join(sec(i) for i in chain(start, fat))[:size]
        mini = b''.join(sec(i) for i in chain(root[2], fat))[:root[3]]
        mfat = []
        for i in chain(mfat_start, fat)[:nmfat]:
            mfat.extend(struct.unpack('<%dI' % per, sec(i)))
        return b''.join(mini[k * mssz:(k + 1) * mssz] for k in chain(start, mfat))[:size]
    raise ValueError('no Workbook stream in the file')


class _Segs:
    """A byte reader over a record plus its CONTINUE records. Character runs that cross into a
    CONTINUE record restart with a new option byte (high-byte flag), as [MS-XLS] 2.5.293 says."""

    def __init__(self, segs):
        self.segs, self.i, self.p = segs, 0, 0

    def _next(self):
        self.i += 1
        self.p = 0
        if self.i >= len(self.segs):
            raise EOFError

    def take(self, n):
        out = b''
        while n > 0:
            if self.p >= len(self.segs[self.i]):
                self._next()
                continue
            k = min(n, len(self.segs[self.i]) - self.p)
            out += self.segs[self.i][self.p:self.p + k]
            self.p += k
            n -= k
        return out

    def chars(self, cch, high):
        out = []
        while cch > 0:
            if self.p >= len(self.segs[self.i]):
                self._next()
                high = self.segs[self.i][0] & 1
                self.p = 1
            w = 2 if high else 1
            k = min(cch, (len(self.segs[self.i]) - self.p) // w)
            if k == 0:
                self.p = len(self.segs[self.i])
                continue
            b = self.segs[self.i][self.p:self.p + k * w]
            out.append(b.decode('utf-16-le' if high else 'latin-1', 'replace'))
            self.p += k * w
            cch -= k
        return ''.join(out)

    def xl_string(self, cch_bytes=2, rich_ok=True):
        """XLUnicodeRichExtendedString (SST) or XLUnicodeString (LABEL, STRING)."""
        cch = struct.unpack('<H' if cch_bytes == 2 else '<B', self.take(cch_bytes))[0]
        flags = self.take(1)[0]
        runs = struct.unpack('<H', self.take(2))[0] if rich_ok and flags & 0x08 else 0
        ext = struct.unpack('<i', self.take(4))[0] if rich_ok and flags & 0x04 else 0
        txt = self.chars(cch, flags & 0x01)
        if runs:
            self.take(4 * runs)
        if ext > 0:
            self.take(ext)
        return txt


def _rk(v):
    """An RK number (MS-XLS 2.5.217): 30-bit integer or the top of a double, maybe x100."""
    if v & 2:
        x = struct.unpack('<i', struct.pack('<I', v))[0] >> 2
    else:
        x = struct.unpack('<d', struct.pack('<Q', (v & 0xFFFFFFFC) << 32))[0]
    return x / 100.0 if v & 1 else x


def _num(x):
    return str(int(x)) if float(x).is_integer() else ('%g' % x)


def read_xls(data):
    """[(sheet name, rows)] from .xls bytes. xlrd when it imports, else the built-in BIFF8 reader."""
    try:
        import xlrd
        wb = xlrd.open_workbook(file_contents=data, on_demand=True)
        out = []
        for sh in wb.sheets():
            rows = []
            for r in range(sh.nrows):
                rows.append([_num(v) if isinstance(v, float) else str(v).strip() for v in sh.row_values(r)])
            out.append((sh.name, rows))
        return out
    except ImportError:
        pass
    wb = _cfb_stream(data)
    recs, pos = [], 0
    while pos + 4 <= len(wb):
        rt, ln = struct.unpack_from('<HH', wb, pos)
        recs.append((pos, rt, wb[pos + 4:pos + 4 + ln]))
        pos += 4 + ln
    biff8 = True
    sheets, sst = [], []
    for k, (pos, rt, d) in enumerate(recs):
        if rt == 0x0809 and k == 0:
            biff8 = struct.unpack_from('<H', d, 0)[0] == 0x0600
        elif rt == 0x0085:
            off, vis, typ = struct.unpack_from('<IBB', d, 0)
            r = _Segs([d[6:]])
            name = r.xl_string(1, False) if biff8 else d[7:7 + d[6]].decode('cp1252', 'replace')
            if typ == 0:
                sheets.append((name, off))
        elif rt == 0x00FC:
            segs, j = [d[8:]], k + 1
            while j < len(recs) and recs[j][1] == 0x003C:
                segs.append(recs[j][2])
                j += 1
            r = _Segs(segs)
            for _ in range(struct.unpack_from('<I', d, 4)[0]):
                try:
                    sst.append(r.xl_string())
                except (EOFError, struct.error):
                    break
        elif rt == 0x000A and sheets:
            break
    index = {p: k for k, (p, _, _) in enumerate(recs)}
    out = []
    for name, off in sheets:
        cells, k = {}, index.get(off)
        if k is None:
            continue
        pending = None
        for pos, rt, d in recs[k + 1:]:
            if rt == 0x000A:
                break
            try:
                if rt == 0x00FD:
                    r, c, _, i = struct.unpack_from('<HHHI', d, 0)
                    cells[(r, c)] = sst[i] if i < len(sst) else ''
                elif rt == 0x0204:
                    r, c = struct.unpack_from('<HH', d, 0)
                    cells[(r, c)] = _Segs([d[6:]]).xl_string(2, False) if biff8 else \
                        d[8:8 + struct.unpack_from('<H', d, 6)[0]].decode('cp1252', 'replace')
                elif rt == 0x0203:
                    r, c, _, x = struct.unpack_from('<HHHd', d, 0)
                    cells[(r, c)] = _num(x)
                elif rt == 0x027E:
                    r, c, _, v = struct.unpack_from('<HHHI', d, 0)
                    cells[(r, c)] = _num(_rk(v))
                elif rt == 0x00BD:
                    r, c0 = struct.unpack_from('<HH', d, 0)
                    for j in range((len(d) - 6) // 6):
                        cells[(r, c0 + j)] = _num(_rk(struct.unpack_from('<I', d, 4 + j * 6 + 2)[0]))
                elif rt == 0x0006:
                    r, c = struct.unpack_from('<HH', d, 0)
                    if d[12:14] == b'\xff\xff':
                        if d[6] == 0:
                            pending = (r, c)
                        elif d[6] == 1:
                            cells[(r, c)] = 'TRUE' if d[8] else 'FALSE'
                    else:
                        cells[(r, c)] = _num(struct.unpack_from('<d', d, 6)[0])
                elif rt == 0x0207 and pending:
                    cells[pending] = _Segs([d]).xl_string(2, False)
                    pending = None
                elif rt == 0x0205:
                    r, c, _, v, err = struct.unpack_from('<HHHBB', d, 0)
                    if not err:
                        cells[(r, c)] = 'TRUE' if v else 'FALSE'
            except (struct.error, EOFError, IndexError):
                continue
        nr = max((r for r, _ in cells), default=-1) + 1
        nc = max((c for _, c in cells), default=-1) + 1
        rows = [[(cells.get((r, c)) or '').strip() for c in range(nc)] for r in range(nr)]
        out.append((name, rows))
    return out



W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _w_text(el):
    parts = []
    for x in el.iter():
        if x.tag == W_NS + 't':
            parts.append(x.text or '')
        elif x.tag == W_NS + 'tab':
            parts.append('\t')
        elif x.tag in (W_NS + 'br', W_NS + 'cr'):
            parts.append('\n')
    return ''.join(parts)


def read_docx(data):
    """(tables, lines) from .docx bytes: tables as [(name, rows)], the body paragraphs as text lines.
    Merged cells (gridSpan) are padded so the columns stay aligned."""
    z = zipfile.ZipFile(io.BytesIO(data))
    body = ET.fromstring(z.read('word/document.xml')).find(W_NS + 'body')
    tables, lines = [], []
    for el in (body if body is not None else []):
        if el.tag == W_NS + 'tbl':
            rows = []
            for tr in el.findall(W_NS + 'tr'):
                cells = []
                for tc in tr.findall(W_NS + 'tc'):
                    cells.append(' '.join(_w_text(p).strip() for p in tc.iter(W_NS + 'p')).strip())
                    span = tc.find('%stcPr/%sgridSpan' % (W_NS, W_NS))
                    if span is not None:
                        cells.extend([''] * (int(span.get(W_NS + 'val', '1')) - 1))
                rows.append(cells)
            tables.append(('Table %d' % (len(tables) + 1), rows))
        elif el.tag == W_NS + 'p':
            lines.extend(_w_text(el).split('\n'))
    return tables, lines



class _TableHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == 'table':
            self.tables.append([])
        elif tag == 'tr' and self.tables:
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = []
            span = dict(attrs).get('colspan')
            self._span = int(span) if span and span.isdigit() else 1

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(re.sub(r'\s+', ' ', ''.join(self.cell)).strip())
            self.row.extend([''] * (self._span - 1))
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.tables[-1].append(self.row)
            self.row = None

    def handle_data(self, d):
        if self.cell is not None:
            self.cell.append(d)


def read_html_tables(text):
    p = _TableHTML()
    p.feed(text)
    return [('Table %d' % (k + 1), t) for k, t in enumerate(p.tables) if t]



def norm_header(s):
    s = re.sub(r'\s+', ' ', s.strip().lower().rstrip(':'))
    if re.fullmatch(r'(?:\w ){2,}\w', s):
        s = s.replace(' ', '')
    return s


def header_map(row):
    """{field: col} for a candidate header row, plus the column where an OUTPUT block starts."""
    cols, cut = {}, None
    for i, cell in enumerate(row):
        h = norm_header(cell)
        if not h:
            continue
        if OUTPUT_HEADER.match(h) and 'name' in cols:
            cut = i
            break
        for f, rx in HEADER_RE:
            if rx.match(h) and f not in cols:
                cols[f] = i
                break
    if 'name' not in cols and 'short' in cols:
        cols['name'], cols['_name_is_short'] = cols['short'], True
    if 'ch' in cols and 'name' in cols and cols['ch'] > cols['name']:
        del cols['ch']
    return cols, cut


def header_windows(row):
    """Column windows [(start, end)] of a header row. Two input tables side by side ("Ch Input Mic | Ch Input Mic")
    give two windows: a new one starts where the first table's first field comes back after a name column. Every
    window must map a name and one more field, else the row is one window."""
    mapped = []
    for i, cell in enumerate(row):
        h = norm_header(cell or '')
        f = next((f for f, rx in HEADER_RE if h and rx.match(h)), None)
        if f:
            mapped.append((i, f))
    if len(mapped) < 4:
        return [(0, len(row))]
    f0, starts = mapped[0][1], [0]
    named = f0 == 'name'
    for i, f in mapped[1:]:
        if f == f0 and named:
            starts.append(i)
            named = f0 == 'name'
        elif f == 'name':
            named = True
    if len(starts) == 1:
        return [(0, len(row))]
    wins = list(zip(starts, starts[1:] + [len(row)]))
    for a, b in wins:
        cols, _ = header_map(row[a:b])
        if 'name' not in cols or len([f for f in cols if not f.startswith('_')]) < 2:
            return [(0, len(row))]
    return wins


def numbered_name(cols, hdr, data):
    """'Input | Source | Mic': a column headed like a name ("Input") that holds channel numbers is the channel column.
    The name is then the next column right of it that holds text: one headed like a name ("Source", "Instrument")
    first, then one headed "Description" or "Notes", then any unmapped one. cols is returned unchanged otherwise."""
    if 'ch' in cols or 'name' not in cols or cols.get('_name_is_short'):
        return cols
    nc = cols['name']
    cells = lambda j: [r[j].strip() for r in data if j < len(r) and r[j] and r[j].strip()]
    vals = cells(nc)
    if len(vals) < 2 or sum(is_int(v) for v in vals) < 0.8 * len(vals):
        return cols
    mapped = {v: k for k, v in cols.items() if not k.startswith('_')}
    best = None
    for j in range(nc + 1, len(hdr)):
        h = norm_header(hdr[j] or '')
        rank = 0 if h and re.match(FIELD_RX['name'], h) and j not in mapped else \
            1 if mapped.get(j) == 'notes' else 2 if j not in mapped else None
        if rank is None:
            continue
        v = cells(j)
        if len(v) < 0.5 * len(vals) or sum(bool(re.search(r'[A-Za-z]', x)) for x in v) < 0.8 * len(v):
            continue
        if best is None or rank < best[0]:
            best = (rank, j)
    if best is None:
        return cols
    out = {k: v for k, v in cols.items() if v != best[1]}
    out['ch'], out['name'] = nc, best[1]
    return out


def find_tables(rows):
    """Candidate (score, header row index, cols, cut, window) for one sheet, best first. cols and cut count from the
    window's first column."""
    cands = []
    for i, row in enumerate(rows[:60]):
        for a, b in header_windows(row):
            cols, cut = header_map(row[a:b])
            cols = numbered_name(cols, row[a:b], [r[a:b] for r in rows[i + 1:i + 80]])
            fields = [f for f in cols if not f.startswith('_')]
            if 'name' not in cols or len(fields) < 2:
                continue
            n = 0
            for r in rows[i + 1:i + 80]:
                r = r[a:b]
                v = r[cols['name']] if cols['name'] < len(r) else ''
                if v and not STOP_NAME.match(v):
                    n += 1
            cands.append((len(fields) * 10 + min(n, 64), i, cols, cut, (a, b)))
    return sorted(cands, key=lambda c: -c[0])


def is_int(s):
    return bool(re.fullmatch(r'\d{1,3}', s.strip()))


def grid_rows(rows, hdr, cols, cut, skipped=None):
    """Input rows under the header row. Rows with a channel number but no name (unused lines on
    the list) go to `skipped` as {'line', 'ch', 'why'} when a list is passed."""
    out, section, blank = [], None, 0
    width = cut if cut else max(len(r) for r in rows) if rows else 0
    name_c = cols['name']
    for k, r in enumerate(rows[hdr + 1:], start=hdr + 2):
        r = (r + [''] * width)[:width] if width else r
        get = lambda f: r[cols[f]].strip() if f in cols and cols[f] < len(r) else ''
        name, ch = get('name'), get('ch')
        nonempty = [c for c in r if c.strip()]
        if not nonempty:
            blank += 1
            if blank >= 4 and out:
                break
            continue
        blank = 0
        if name and STOP_NAME.match(name):
            break
        if len(nonempty) == 1 and not is_int(nonempty[0]) and not ch:
            section = nonempty[0]
            continue
        if not name:
            m = re.fullmatch(r'(?i)\s*(?:ch(?:annel)?\.?\s*|#\s*)?(\d{1,3})\s*', ch)
            if m and skipped is not None:
                skipped.append({'line': k, 'ch': int(m.group(1)), 'why': 'blank'})
            continue
        chm = re.fullmatch(r'(?i)\s*(?:ch(?:annel)?\.?\s*|#\s*|in\s*)?(\d{1,3})(?:\s*[-–/&+]\s*(\d{1,3}))?\s*', ch)
        row = {'line': k, 'ch': int(chm.group(1)) if chm else None, 'name': name,
               'mic': get('mic'), 'stand': get('stand'), 'p48': p48_cell(get('p48')), 'notes': get('notes'),
               'who': get('who'), 'patch': get('patch'), 'section': section, 'raw': ' | '.join(nonempty)}
        if chm and chm.group(2) and int(chm.group(2)) == int(chm.group(1)) + 1:
            row['_range'] = int(chm.group(2))
        prev = out[-1] if out else None
        if ch and prev is not None and prev.get('_ch_text') == ch:
            row['ch'] = prev.pop('_range', None)
            row.pop('_range', None)
        row['_ch_text'] = ch
        if STEREO_CELL.match(get('stereo')) or STEREO_NOTE.search(row['notes']):
            row['_stereo_hint'] = True
        elif MIC_PAIR.search(row['mic']):
            row['_mic_pair'] = True
        if re.match(r'(?i)^(opt|option|alt|spare\?)', ch):
            row['option'] = True
        if cols.get('_name_is_short'):
            row['short'] = name
        elif 'short' in cols:
            row['short'] = get('short')
        if row['p48'] is None:
            row['p48'] = p48_text(' '.join([row['mic'], row['notes']]))
        out.append(row)
    return out


def p48_cell(v):
    v = v.strip().lower()
    if not v:
        return None
    if v in ('x', '*', 'y', 'yes', '1', 'true', 'on', '48', '48v', '+48', '+48v', '✓', '✔', 'p48', 'ph'):
        return True
    if v in ('-', 'n', 'no', '0', 'false', 'off', 'none', 'no 48v'):
        return False
    t = p48_text(v)
    if t is None and v.split()[0] in ('x', '*', 'y', 'yes', '✓', '✔', '48v', '+48'):
        return True
    return t


def p48_text(s):
    if NO48_RX.search(s):
        return False
    return True if P48_RX.search(s) else None



SECTION_WORD = re.compile(r'(?i)\b(drums?|kit|band|vocals|vox|keys|keyboards|guitars|strings|horns|brass|winds|perc\w*|'
                          r'tracks|playback|fx|effects|returns|misc\w*|utility|comms|stage|sl|sr|usc|dsc|sub|riser|choir|'
                          r'rhythm|inputs?|wireless|rf|mics|lines|audience|ambient|monitors?|foh|other|extras?|section|'
                          r'rig|world|front|back|side|left|right|center|centre)\b')
LEAD_NUM = re.compile(r'^\s*(?:ch(?:annel)?\.?\s*|#\s*|in\s*)?(\d{1,3})(?:\s*[-–/&+]\s*(\d{1,3}))?(?:\s*[.):\-–]\s*|\s+|$)', re.I)
NOTE_LINE = re.compile(r"(?i)^\s*(?:notes?\b|nb\b|n\.b\.|ps\b|p\.s\.|please\b|pls\b|thanks?\b|thank you|cheers\b|"
                       r"contact\b|questions?\b|totals?\b|all (?:mics|inputs|channels)\b|\*)")
BULLET = re.compile(r'^\s*(?:[-*•·▪◦>]+|\(\w\)|\w\))\s+')
SECTION_LINE = re.compile(r'^\s*(?:#+\s*)?([A-Za-z][\w &/+\'’.-]{0,40}?)\s*:\s*$')


def split_cells(s):
    """Split a text row on strong delimiters; returns list of cells."""
    for d in ('\t', '|', ';'):
        if d in s:
            return [c.strip() for c in s.split(d) if c.strip()]
    if re.search(r'\S {2,}\S', s):
        return [c.strip() for c in re.split(r' {2,}', s.strip()) if c.strip()]
    return [s.strip()]


def classify_cell(c):
    if P48_RX.fullmatch(c.strip()) or c.strip().lower() in ('x', '*', '48', 'p'):
        return 'p48'
    if STAND_RX.fullmatch(c.strip()):
        return 'stand'
    if MIC_RX.search(c) and len(c.split()) <= 6:
        return 'mic'
    if re.fullmatch(r"\d{1,3}\s?(?:'|ft|feet|m)", c.strip()):
        return 'xlr'
    return None


def parse_text_line(body):
    """name/mic/stand/p48/notes from the text after the channel number."""
    row = {'name': '', 'mic': '', 'stand': '', 'p48': None, 'notes': ''}
    body = body.strip()
    notes = re.findall(r'\(([^)]*)\)', body)
    p = p48_text(body)
    row['p48'] = p
    full_body = body
    body = re.sub(r'\s*\([^)]*\)\s*', ' ', body).strip()
    if '(' in body:
        body = body[:body.index('(')].strip()
        notes.append(full_body[full_body.index('(') + 1:])
    cells = split_cells(body)
    if len(cells) == 1:
        for sep in (' - ', ' – ', ', ', ' / '):
            if sep in body:
                left, right = body.split(sep, 1)
                if classify_cell(right) or MIC_RX.search(right) or STAND_RX.search(right) or P48_RX.search(right):
                    cells = [left] + [c.strip() for c in re.split(r',\s*| - | – ', right) if c.strip()]
                    break
    if len(cells) > 1:
        rest = []
        for c in cells:
            k = classify_cell(c) if row['name'] else None
            if not row['name']:
                row['name'] = c
            elif k == 'p48':
                continue
            elif k == 'stand' and not row['stand']:
                row['stand'] = c
            elif k == 'mic' and not row['mic']:
                row['mic'] = c
            elif k == 'xlr':
                continue
            else:
                rest.append(c)
        row['notes'] = '; '.join(rest)
    else:
        txt = re.sub(r'\([^)]*\)', ' ', body)
        txt = P48_RX.sub(' ', txt)
        stand = STAND_RX.search(txt)
        if stand and stand.group(0).strip().lower() != 'stand':
            row['stand'] = stand.group(0).strip()
            txt = txt[:stand.start()] + ' ' + txt[stand.end():]
        m = None
        for mm in MIC_RX.finditer(txt):
            if mm.start() == 0:
                continue
            if re.fullmatch(r'(?i)d\.?i\.?|line|rf|wireless|xlr|direct', mm.group(0)) and not txt[mm.end():].strip():
                row['mic'] = mm.group(0)
                break
            m = mm
            break
        if m:
            row['name'] = txt[:m.start()].strip(' -–,:/')
            row['mic'] = txt[m.start():].strip(' -–,:/')
        else:
            row['name'] = txt.strip(' -–,:/')
    row['name'] = re.sub(r'\s+', ' ', P48_RX.sub('', row['name'])).strip(' -–,:/')
    row['name'] = re.sub(r'(?i)\s+\(?mono\)?$', '', row['name'])
    m = re.match(r'(?i)^(?:di|di line|line|mic|xlr|input|line in)\s*:\s*(.+)$', body.strip())
    if m:
        row['name'], row['mic'] = m.group(1).strip(' .'), body.split(':')[0].strip()
        row['_described'] = True
    keep = []
    for n in notes:
        if MIC_RX.search(n) and len(n.split()) <= 4 and not P48_RX.fullmatch(n.strip()) and \
                (not row['mic'] or GENERIC_RX.match(row['mic'])):
            row['mic'] = n.strip()
        elif n.strip() and not P48_RX.fullmatch(n.strip()):
            keep.append(n.strip())
    if keep:
        row['notes'] = '; '.join(keep + ([row['notes']] if row['notes'] else []))
    return row


PROSE_RX = re.compile(r"(?i)\b(?:i'?ll|i'?m|i|we|we'?ll|you|need|needs|want|would|could|should|maybe|potentially|probably|"
                      r"please|bring|using|plays?|playing|sings?|singing|via|also|then|with a|on a|our|my|his|her|their)\b")


GREETING = re.compile(r"(?i)^(?:hey|hi|hello|thanks|thank you|cheers|best|regards|sent from|here'?s|see you|let me know)\b[^\n]{0,60}[!.:]?$")


def is_prose(t, r):
    """A sentence rather than an input-list row: send it to the model fallback."""
    words = t.split()
    hits = len(PROSE_RX.findall(t))
    return (len(words) > 7 and not r['mic']) or hits >= 2 or (hits >= 1 and len(words) >= 5)


def _numbered(line):
    s2 = BULLET.sub('', line)
    m = LEAD_NUM.match(s2)
    return bool(m and int(m.group(1)) > 0 and not re.match(r'^\s*\d{1,3}\s*(?:x|pcs?)\b', s2, re.I))


def text_rows(text, pdf=False, skipped=None, dropped=None):
    """Rows from pasted text or a PDF's text layer. In a numbered list (most lines start with a channel number) an
    unnumbered short line followed by a numbered one is a section heading ("BASS", "GUITAR"), and an unnumbered
    note ("Note: 48V on 3, 7", "Please ...") is left out; those go to `dropped` as {'line', 'text'}."""
    lines = text.replace('\r', '').split('\n')
    if pdf:
        lines = _pdf_columns(lines)
    out, section, auto = [], None, 0
    seen_numbers = False
    filled = [l for l in lines if l.strip()]
    num_list = not pdf and sum(map(_numbered, filled)) >= max(3, 0.6 * len(filled))
    for k, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        s = line.rstrip()
        m = SECTION_LINE.match(s)
        if m and not LEAD_NUM.match(s):
            section = m.group(1).strip()
            continue
        if num_list and not _numbered(s):
            t = BULLET.sub('', s).strip()
            later = next((l for l in lines[k:] if l.strip()), '')
            if GREETING.match(t) and categorize(t)[0] == 'other':
                continue
            if NOTE_LINE.match(t):
                if dropped is not None:
                    dropped.append({'line': k, 'text': t[:80]})
                continue
            if _numbered(later) and len(t.split()) <= 4 and not re.search(r'\d', t) and not MIC_RX.search(t) and \
                    (t.upper() == t or t.endswith(':')):
                section = t.rstrip(':').strip().title()
                continue
        inline = re.match(r'^\s*([A-Za-z][\w &/+-]{1,24}):\s+(.+)$', s)
        s2 = BULLET.sub('', s)
        num = LEAD_NUM.match(s2)
        if num and int(num.group(1)) > 0 and not re.match(r'^\s*\d{1,3}\s*(?:x|pcs?)\b', s2, re.I):
            seen_numbers = True
            a, b = int(num.group(1)), num.group(2)
            body = s2[num.end():]
            if not body.strip() or not re.search(r'[A-Za-z0-9]', body):
                if skipped is not None:
                    skipped.append({'line': k, 'ch': a, 'why': 'blank'})
                continue
            r = parse_text_line(body)
            r.update(line=k, ch=a, section=section, raw=s.strip(), who='', patch='', _text=True)
            if b and int(b) == a + 1:
                r['_range'] = int(b)
            out.append(r)
            continue
        if pdf:
            t = s.strip()
            if len(t) <= 40 and t.upper() == t and re.search(r'[A-Z]{3}', t) and not MIC_RX.search(t):
                section = t.title()
            continue
        if inline and not BULLET.match(s) and ',' in inline.group(2):
            section = inline.group(1).strip()
            for item in re.split(r',\s*|\s+and\s+', inline.group(2)):
                if item.strip():
                    r = parse_text_line(item)
                    r.update(line=k, ch=None, section=section, raw=s.strip(), who='', patch='')
                    r['_free'] = True
                    out.append(r)
            continue
        t = s2.strip()
        if not t or (GREETING.match(t) and categorize(t)[0] == 'other'):
            continue
        if t.upper() == t and len(t) <= 30 and not MIC_RX.search(t) and len(t.split()) <= 4 and not BULLET.match(s) \
                and SECTION_WORD.search(t):
            section = t.title()
            continue
        r = parse_text_line(t)
        r.update(line=k, ch=None, section=section, raw=s.strip(), who='', patch='', _text=True)
        if is_prose(t, r):
            r['_prose'] = True
        out.append(r)
    return out


def _cells_pos(line):
    """[(start, end, text)] for runs separated by 2+ spaces."""
    return [(m.start(), m.end(), m.group(0)) for m in re.finditer(r'\S+(?: \S+)*', line)]


def _delspace(t):
    """"D R U M P L AY B A C K" -> "DRUM PLAYBACK"-ish: collapse letter-spaced headings."""
    if re.fullmatch(r'(?:\S{1,2} )+\S{1,2}', t.strip()):
        return re.sub(r'(?<=\S) (?=\S)', '', t.strip())
    return t.strip()


def layout_rows(lines):
    """Rows from fixed-width text (pdftotext -layout) that has a header line. Each cell goes to the
    header column whose span holds the cell's centre. Returns None when no header is found."""
    out, cols, section, found = [], None, None, False
    for k, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        cells = _cells_pos(line)
        hm, _ = header_map([c[2] for c in cells])
        fields = [f for f in hm if not f.startswith('_')]
        if 'name' in hm and len(fields) >= 3:
            centres = [((s + e) / 2.0, i) for i, (s, e, t) in enumerate(cells)]
            fmap = {i: f for f, i in hm.items() if not f.startswith('_')}
            cols = [(c, fmap.get(i)) for c, i in centres]
            found = True
            continue
        if not cols:
            continue
        t = line.strip()
        if not LEAD_NUM.match(t):
            if len(t) <= 48 and not MIC_RX.search(_delspace(t)) and re.search(r'[A-Za-z]{2}', t):
                h = _delspace(t)
                if h.upper() == h or len(h.split()) <= 3:
                    section = h.title()
            continue
        bounds = [(cols[i][0] + cols[i + 1][0]) / 2.0 for i in range(len(cols) - 1)]
        got = {}
        for s, e, txt in cells:
            c = (s + e) / 2.0
            j = 0
            while j < len(bounds) and c > bounds[j]:
                j += 1
            f = cols[j][1]
            if f:
                got[f] = (got[f] + ' ' + txt) if f in got else txt
        num = re.match(r'\s*(\d{1,3})', got.get('ch', '') or t)
        name = got.get('name', '').strip()
        if not num or not name:
            continue
        row = {'line': k, 'ch': int(num.group(1)), 'name': name, 'mic': got.get('mic', ''), 'stand': got.get('stand', ''),
               'p48': p48_cell(got.get('p48', '')), 'notes': got.get('notes', ''), 'who': got.get('who', ''),
               'patch': got.get('patch', ''), 'section': section, 'raw': t}
        if row['p48'] is None:
            row['p48'] = p48_text(row['mic'] + ' ' + row['notes'])
        out.append(row)
    return out if found and out else None


def _pdf_columns(lines):
    """pdftotext -layout keeps side-by-side tables on one line. Split each line at the x position
    where a second "<num>  <Name>" column starts, then emit the left column first, page by page."""
    pages, cur = [], []
    for l in lines:
        if '\f' in l:
            parts = l.split('\f')
            cur.append(parts[0]); pages.append(cur); cur = [parts[-1]]
        else:
            cur.append(l)
    pages.append(cur)
    out = []
    for pg in pages:
        split = None
        for l in pg:
            cells = _cells_pos(l)
            names = [norm_header(t) for _, _, t in cells]
            idx = [i for i, h in enumerate(names) if re.match(FIELD_RX['ch'], h)]
            nidx = [i for i, h in enumerate(names) if re.match(FIELD_RX['name'], h)]
            if len(idx) >= 2 and len(nidx) >= 2:
                split = cells[idx[1]][0] - 2
                break
        if not split:
            out.extend(pg)
            continue
        left, right = [], []
        for l in pg:
            if len(l) > split and l[:split].strip() and l[split:].strip():
                left.append(l[:split]); right.append(l[split:])
            elif len(l) > split and not l[:split].strip():
                right.append(l[split:])
            else:
                left.append(l)
        out.extend(left + right)
    return out



def expand_pairs(rows):
    """"OH L/R", "Keys (stereo)", "8-9 OH", "2x OH", "Keys - stereo DI" -> two rows. So do a Stereo column, "stereo"
    in the notes and a mic pair ("KM184 x2") on a source that is a pair by nature.
    A range row that names one side ("3-4  OH L") is the L half of a pair whose R row follows with its channel cell
    blank (or merged): it stays one row and lends the R row the second number. A range row followed by an unnumbered
    row of the same name ("3-4 OH" then "OH") is one pair: the second row is dropped."""
    out, i = [], 0
    rows = list(rows)
    while i < len(rows):
        r = rows[i]
        i += 1
        name = r['name']
        nxt = rows[i] if i < len(rows) else None
        if r.get('_range') and side_of(name)[0]:
            r = dict(r)
            end = r.pop('_range')
            if nxt is not None and nxt.get('ch') is None and side_of(nxt['name'])[0]:
                rows[i] = dict(nxt, ch=end)
            out.append(r)
            continue
        if r.get('_range') and nxt is not None and nxt.get('ch') is None and not side_of(nxt['name'])[0] and \
                _base_key(nxt['name']) == _base_key(name):
            i += 1
        qty = re.match(r'(?i)^\s*([2-8])\s*x\s+(.+)$', name)
        if qty:
            name = qty.group(2).strip()
            r = dict(r, name=name)
        stereo_mic = bool(r.get('_text') and r.get('mic') and PAIR_WORDS.search(r['mic']) and ',' not in r['mic']
                          and not side_of(name)[0])
        hinted = not side_of(name)[0] and (r.get('_stereo_hint') or (
            r.get('_mic_pair') and categorize(name, r.get('mic', ''), r.get('section'))[0] in STEREO_NATURE))
        if PAIR_WORDS.search(name) or r.get('_range') or stereo_mic or hinted or (qty and qty.group(1) == '2'):
            base = re.sub(r'\s+', ' ', PAIR_WORDS.sub(' ', name)).strip()
            mic = PAIR_WORDS.sub(' ', r.get('mic') or '').strip() if stereo_mic else r.get('mic')
            if hinted and r.get('_mic_pair'):
                mic = re.sub(r'\s+', ' ', MIC_PAIR.sub(' ', mic or '')).strip(' ,')
            if base:
                a = dict(r, name=base + ' L', mic=mic); b = dict(r, name=base + ' R', mic=mic)
                if r['ch'] is not None:
                    b['ch'] = r['ch'] + 1
                a.pop('_range', None); b.pop('_range', None)
                a['_expanded'] = b['_expanded'] = True
                out += [a, b]
                continue
        if qty:
            for k in range(int(qty.group(1))):
                out.append(dict(r, name='%s %d' % (name, k + 1), _expanded=True,
                                ch=None if r['ch'] is None else r['ch'] + k))
            continue
        out.append(r)
    return out


def side_of(name):
    """('L'|'R'|None, base name without side tokens, stage_side flag). Every standalone side token in the
    name must agree ("SPD Master L - Trax L" is L; "SPD Sub L - Hits" is L too)."""
    m = re.search(r'(?i)\bstage (left|right)\b', name)
    if m:
        return ('L' if m.group(1).lower() == 'left' else 'R'), (name[:m.start()] + name[m.end():]).strip(), True
    toks = name.split()
    sides = [(i, SIDE_TOKENS[t.lower().strip(',.:;')]) for i, t in enumerate(toks) if t.lower().strip(',.:;') in SIDE_TOKENS]
    if sides and len({s for _, s in sides}) == 1:
        drop = {i for i, _ in sides}
        return sides[0][1], ' '.join(t for i, t in enumerate(toks) if i not in drop), False
    if sides:
        return None, name, False
    for i in range(len(toks) - 1, -1, -1):
        t = toks[i].lower().strip(',.')
        if t in STAGE_SIDE:
            return STAGE_SIDE[t], ' '.join(toks[:i] + toks[i + 1:]), True
    return None, name, False


def _base_key(base):
    b = base.lower()
    b = re.sub(r'\b\d+\b', '#', b)
    return re.sub(r'[^a-z#]+', ' ', b).strip()


def norm_name(s):
    return re.sub(r'\s+', ' ', (s or '').strip().lower())


def pair_rows(rows, stereo=True):
    """Group L/R neighbours into one stereo channel. Returns list of lists of rows."""
    groups, i = [], 0
    while i < len(rows):
        r = rows[i]
        if stereo and i + 1 < len(rows):
            n = rows[i + 1]
            s1, b1, st1 = side_of(r['name'])
            s2, b2, st2 = side_of(n['name'])
            if not s1 and not s2 and norm_name(r['name']) == norm_name(n['name']) and r.get('mic') and n.get('mic'):
                s1, b1, st1 = side_of(r['mic'])[0], r['name'], False
                s2, b2, st2 = side_of(n['mic'])[0], n['name'], False
            ok = s1 and s2 and s1 != s2 and _base_key(b1) == _base_key(b2) and _base_key(b1)
            if not ok and r.get('short') and n.get('short'):
                t1, t2 = side_of(r['short']), side_of(n['short'])
                if t1[0] and t2[0] and t1[0] != t2[0] and _base_key(t1[1]) == _base_key(t2[1]) and _base_key(t1[1]):
                    (s1, b1, st1), (s2, b2, st2), ok = t1, t2, True
            if ok and (st1 or st2):
                cat = categorize(b1, r.get('mic', ''), r.get('section'))[0]
                ok = cat in ('overhead', 'audience', 'drum_room', 'choir', 'playback', 'music_playback', 'keys')
            if ok and r.get('ch') is not None and n.get('ch') is not None:
                ok = n['ch'] == r['ch'] + 1
            if ok:
                groups.append([r, n] if s1 == 'L' else [n, r])
                i += 2
                continue
        groups.append([r])
        i += 1
    return groups


def _vocal_kind(n, mic, artist, p):
    """lead_vocal or backing_vocal for a vocal channel."""
    if artist and any(len(a) > 2 and re.search(r'\b%s\b' % re.escape(a.lower()), n) for a in artist.lower().split()):
        return 'lead_vocal', 0.8
    spare = re.search(r'spare|spr|guest|gst', n)
    if re.search(r'\bmain\b', n) and not spare:
        return 'lead_vocal', 0.7
    if re.search(r'\b(rf|wireless|axient|ad ?[12]|ulx|qlx)\b', (mic or '').lower() + ' ' + n) and not spare:
        return 'lead_vocal', 0.6
    return 'backing_vocal', p


def categorize(name, mic='', section=None, artist=None):
    """(category, confidence) from the channel name, then the mic, then the section heading."""
    n = ' ' + re.sub(r'[_]+', ' ', (name or '').lower().strip()) + ' '
    if re.match(r'^\s*(?:t\.|t\s|(?:trax|trx|tracks?|pb\s?\d+|playback)(?![a-z]))', n) and \
            not re.search(r'click|clk|guide|cue|smpte|ltc|time ?code|midi|tb|talk', n):
        return 'playback', 0.9
    for cat, rx, p in CAT_RULES:
        if rx.search(n.strip()) if cat == 'spare' else rx.search(n):
            if cat == 'vocal':
                return _vocal_kind(n, mic, artist, p)
            if cat == 'drum_pad' and re.search(r'(?i)ableton|laptop|computer|mac|tracks|playback|mainstage|loop', mic or ''):
                return 'playback', 0.6
            return cat, p
    for rx, cat in MIC_PRIOR:
        if mic and rx.search(mic):
            if cat == 'vocal':
                k, _ = _vocal_kind(n, mic, artist, 0.55)
                return k, 0.55
            return cat, 0.55
    if section:
        for rx, cat in SECTION_PRIOR:
            if rx.search(section):
                return cat, 0.45
    return 'other', 0.2


def is_person(tok):
    """A capitalised word that is not audio vocabulary: a player's name ("Riley", "Alex")."""
    return bool(tok) and tok[:1].isupper() and tok[1:2].islower() and tok.isalpha() and len(tok) >= 3 and \
        tok.lower() not in NOT_PERSON and tok.lower() not in WORDS and \
        not any(rx.search(' %s ' % tok.lower()) for c, rx, p in CAT_RULES if c != 'spare')


def abbreviate(full, stereo=False, max_len=MAX_NAME, initials=False):
    """An 8-character console name in a common touring style ("Snr1 Top", "G.Nord", "T. Keys").
    initials=True writes every "<Player> <thing>" as "P.thing" even when it fits ("Lee TB" -> "L.TB"),
    a common habit on dLive names; the artist's own vocal is the caller's exception."""
    s = re.sub(r'\([^)]*\)', ' ', full or '')
    s = DROP_WORDS.sub(' ', s)
    s = re.sub(r'\s+', ' ', s).strip(' -–:/,')
    if stereo:
        side, base, _ = side_of(s)
        if side:
            s = base
        s = re.sub(r'(?i)\s*\b(?:l ?/ ?r|stereo|st)\b', '', s).strip()
        s = re.sub(r'(?i)(?<=\w) (?:l|r)\b(?= -|$)', '', s).strip()
    s = re.sub(r'\s+', ' ', s).strip(' -–:/,')
    if not s:
        return ''
    if initials:
        t = s.split()
        if len(t) >= 2 and is_person(t[0]):
            return t[0][0] + '.' + abbreviate(' '.join(t[1:]), False, max_len - 2)
    if len(s) <= max_len:
        return s
    m = re.match(r'(?i)^(trax\.?|trx\.?|tracks?|pb|playback|t\.?)\s*\d*\s*[-–—:]\s*(.+)$', s)
    if m:
        return 'T.' + abbreviate(m.group(2), False, max_len - 2)
    m = re.match(r'(?i)^(spd|pad)\b[^-–]*[-–—]\s*(.+)$', s)
    if m:
        return 'SPD.' + abbreviate(m.group(2), False, max_len - 4)
    if ' / ' in s and len(s) > max_len:
        alts = [a.strip() for a in s.split(' / ') if a.strip()]
        fits = [a for a in alts if len(abbreviate(a, False, 99)) <= max_len or
                len(re.sub(r' ', '', abbreviate(a, False, 99))) <= max_len]
        s = fits[-1] if fits else alts[0]
    s = s.replace(' - ', ' ').replace(' – ', ' ').replace(' — ', ' ')
    if len(s) <= max_len:
        return s
    for rx, rep in PHRASES:
        s = rx.sub(rep, s)
    s = re.sub(r'\s+', ' ', s).strip()
    if len(s) <= max_len:
        return s
    toks = s.split()
    first = toks[0]
    if len(toks) >= 2 and toks[-1].lower() in ('rf', 'di', 'wireless', 'wired', 'mic', 'xlr'):
        return abbreviate(' '.join(toks[:-1]), False, max_len)
    toks = [WORDS.get(t.lower(), t) for t in toks]
    toks = [t for t in toks if t]
    s = ' '.join(toks)
    if len(s) <= max_len:
        return s
    if len(toks) >= 2 and toks[0] == first and is_person(first):
        return first[0] + '.' + abbreviate(' '.join(toks[1:]), False, max_len - 2)
    s = re.sub(r'(?<=[A-Za-z.]) (?=\d)', '', s)
    if len(s) <= max_len:
        return s
    toks = s.split()
    joined = ''.join(t[:1].upper() + t[1:] if i and t[:1].isalpha() else t for i, t in enumerate(toks))
    if len(joined) <= max_len:
        return joined
    while len(''.join(toks)) > max_len:
        k = max(range(len(toks)), key=lambda j: len(toks[j]) - (2 if '.' in toks[j][:3] else 0))
        w = toks[k]
        pre, w = (w[:2], w[2:]) if re.match(r'^[A-Za-z]\.', w) else ('', w)
        nw = pre + w[:1] + re.sub(r'(?i)[aeiou]', '', w[1:])
        w = pre + w
        if nw == w or len(w) <= 3:
            break
        toks[k] = nw
        if len(' '.join(toks)) <= max_len:
            return ' '.join(toks)
    s = ''.join(t[:1].upper() + t[1:] if i and t[:1].isalpha() else t for i, t in enumerate(toks))
    return s[:max_len]


def tidy_case(name):
    """"kick out" -> "Kick Out" when the whole name was typed in lower case."""
    if name and name == name.lower() and re.search('[a-z]', name):
        return ' '.join(w[:1].upper() + w[1:] if w.isalpha() and len(w) > 2 or w == 'in' else
                        (w.upper() if w in ('oh', 'di', 'tb', 'ag', 'eg', 'hh', 'lv', 'bgv') else w) for w in name.split())
    return name


def uniquify(names):
    seen, out = {}, []
    for n in names:
        k = n.lower()
        if k in seen and n:
            seen[k] += 1
            suf = str(seen[k])
            n = (n[:MAX_NAME - len(suf)].rstrip() + suf)
        else:
            seen[k] = 1
        out.append(n)
    return out


def source_type(mic, name):
    t = (mic or '').lower() + ' ' + (name or '').lower()
    if re.search(r'\b(rf|wireless|axient|ad ?[12]|ulx|qlx|ur ?4d|ew ?\d{3}|psm)\b', t):
        return 'rf'
    if re.search(r'\b(madi|dante|aes|usb|avb|optocore)\b', t):
        return 'digital'
    if re.search(r'\b(d\.?i\.?|j ?48|jdi|rndi|livewire|pro ?48|pro ?d2|ldi|usb ?di|reamp)\b', t):
        return 'di'
    if re.search(r'\b(line|xlr|trs|1/4)\b', t) and not MIC_RX.search(re.sub(r'\b(line|xlr|trs)\b', '', mic or '')):
        return 'line'
    if mic and MIC_RX.search(mic):
        return 'mic'
    return None


def build_channels(rows, stereo=True, artist=None, initials=True):
    rows = expand_pairs(rows)
    channels, unparsed = [], []
    for g in pair_rows(rows, stereo):
        a = g[0]
        st = len(g) == 2
        full = a['name']
        if st:
            side, base, _ = side_of(a['name'])
            full = base
        cat, cconf = categorize(full, a.get('mic', ''), a.get('section'), artist)
        mic = a.get('mic') or (g[1].get('mic') if st else '')
        p48 = a.get('p48')
        if st and p48 is None:
            p48 = g[1].get('p48')
        flags, conf = [], 0.95
        if a.get('_prose'):
            flags.append('prose_line'); conf = 0.3
        if a.get('_free'):
            flags.append('free_text_item'); conf = min(conf, 0.6)
        if a.get('_described'):
            flags.append('described_line'); conf = min(conf, 0.5)
        if not full:
            flags.append('no_name'); conf = 0.2
        if cconf < 0.6:
            flags.append('category_guess')
        if a.get('_expanded'):
            flags.append('expanded_from_pair')
        if a.get('option'):
            flags.append('option')
        if st and side_of(a['name'])[2]:
            flags.append('paired_on_SL_SR')
        short = a.get('short') if a.get('short') and len(a.get('short', '')) <= MAX_NAME else None
        if short and st and side_of(short)[0] and side_of(short)[1]:
            short = side_of(short)[1]
        name = short or tidy_case(abbreviate(a['name'], st, initials=initials and cat != 'lead_vocal'))
        if not short and len(abbreviate(full, False, 99)) > MAX_NAME:
            flags.append('name_shortened')
        ch = {
            'number': None,
            'name': name,
            'full_name': full,
            'source': {'mic': mic or None, 'type': source_type(mic, full),
                       'patch': [r.get('patch') or None for r in g] if any(r.get('patch') for r in g) else None},
            'stereo': st,
            'list_channels': [r.get('ch') for r in g],
            'phantom': p48,
            'category': cat,
            'category_confidence': round(cconf, 2),
            'who': a.get('who') or None,
            'stand': a.get('stand') or None,
            'notes': '; '.join(x for x in (a.get('notes'), g[1].get('notes') if st else None) if x) or None,
            'section': a.get('section'),
            'confidence': round(conf, 2),
            'flags': flags,
            'raw': [r.get('raw') for r in g],
        }
        if 'prose_line' in flags or conf < 0.6:
            unparsed.append({'line': a.get('line'), 'text': a.get('raw'), 'channel': len(channels) + 1})
        channels.append(ch)
    for n, (c, nm) in enumerate(zip(channels, uniquify([c['name'] for c in channels])), start=1):
        c['number'] = n
        c['name'] = nm
    return channels, unparsed



def make_request(rows, meta, stereo=True, artist=None, model=None, firmware=None, rate=None, intent=None, naming='initials'):
    channels, unparsed = build_channels(rows, stereo, artist, initials=(naming == 'initials'))
    return {
        'schema': 'digico-session-request',
        'version': SCHEMA_VERSION,
        'console': {'model': model, 'firmware': firmware, 'sample_rate': rate},
        'show': {'artist': artist},
        'channels': channels,
        'buses': None,
        'fx': None,
        'intent': intent,
        'options': {'stereo_pairs': 'auto' if stereo else 'mono', 'naming': naming,
                    'processing': 'starting_points', 'include_options': True},
        'parse': dict(meta, rows=len(rows), channels=len(channels),
                      stereo=sum(c['stereo'] for c in channels), unparsed=unparsed),
    }


def _col_letters(i):
    s = ''
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


class _Block:
    """One input table: a sheet (or Word table), a header row and a column window. rows() parses it once."""

    def __init__(self, score, sheet, si, rows, hdr, cols, cut, win):
        self.score, self.sheet, self.si, self.all_rows, self.hdr = score, sheet, si, rows, hdr
        self.cols, self.cut, self.win = cols, cut, win
        self._parsed = None

    def rows(self):
        if self._parsed is None:
            a, b = self.win
            skipped = []
            got = grid_rows([r[a:b] for r in self.all_rows], self.hdr, self.cols, self.cut, skipped)
            for r in got:
                r['sheet'] = self.sheet
            self._parsed = (got, skipped)
        return self._parsed

    def numbers(self):
        """Channel numbers on this block (named rows, ranges and numbered blank lines), or [] when under 60% of its
        rows carry one."""
        got, skipped = self.rows()
        if not got or sum(r['ch'] is not None for r in got) < 0.6 * len(got):
            return []
        nums = [r['ch'] for r in got if r['ch'] is not None] + [r['_range'] for r in got if r.get('_range')]
        return nums + [s['ch'] for s in skipped]

    def label(self):
        a, b = self.win
        width = max((len(r) for r in self.all_rows), default=0)
        cols = '' if a == 0 and b >= width else ' columns %s-%s' % (_col_letters(a), _col_letters(min(b, width) - 1))
        nums = self.numbers()
        return '%s%s%s' % (self.sheet, cols, ' (ch %d-%d)' % (min(nums), max(nums)) if nums else '')


CONTINUE_GAP = 8


def parse_grid_sheets(sheets, want=None):
    """(rows, meta) from the best input table on any sheet (or Word table), plus the blocks that carry its channel
    numbering on: the right half of a list printed side by side, a second sheet "Inputs 13-22", a second Word table
    on page 2 (with or without its own header row). Blocks that look like input tables but were not used are listed
    in meta['unused']."""
    blocks, alts, bare = [], [], []
    years = [int(y) for s, _ in sheets for y in re.findall(r'\b(20\d\d)\b', s or '')]
    year_max = max(years) if years else None
    for si, (name, rows) in enumerate(sheets):
        if want and name != want:
            continue
        cands = find_tables(rows)
        if not cands:
            bare.append((si, name, rows))
            continue
        bonus = 15 if re.search(r'(?i)\bio\b|input', name or '') else 0
        if year_max and str(year_max) in (name or ''):
            bonus += 12
        bonus -= 25 if re.search(r'(?i)output|loom|plot|network|count', name or '') else 0
        hdr = cands[0][1]
        alts.append({'sheet': name, 'score': cands[0][0] + bonus, 'header_row': hdr + 1})
        for score, h, cols, cut, win in cands:
            if h == hdr:
                blocks.append(_Block(score + bonus, name, si, rows, h, cols, cut, win))
    if not blocks:
        return None, {'error': 'no input table found', 'sheets': [s for s, _ in sheets]}
    best = max(blocks, key=lambda b: (b.score, -b.si, -b.win[0]))
    keys = {f for f in best.cols if not f.startswith('_')}
    for si, name, rows in bare:
        width = max((len(r) for r in rows), default=0)
        if width > max(v for k, v in best.cols.items() if not k.startswith('_')):
            blocks.append(_Block(0, name, si, rows, -1, best.cols, best.cut, (0, width)))
    taken = [best]
    nums = best.numbers()
    if nums and 'ch' in keys:
        lo, hi = min(nums), max(nums)
        grew = True
        while grew:
            grew = False
            for b in blocks:
                if b in taken or 'ch' not in b.cols or not b.numbers():
                    continue
                n = b.numbers()
                if hi < min(n) <= hi + CONTINUE_GAP:
                    taken.append(b)
                    hi, grew = max(n), True
                elif lo - CONTINUE_GAP <= max(n) < lo:
                    taken.insert(0, b)
                    lo, grew = min(n), True
    rows_out, skipped = [], []
    for b in taken:
        got, sk = b.rows()
        rows_out += got
        skipped += sk
    unused = [b for b in blocks if b not in taken and b.hdr >= 0 and b.score >= 0.6 * best.score and
              len(b.rows()[0]) >= 3]
    meta = {'sheet': best.sheet, 'header_row': best.hdr + 1,
            'columns': {f: c for f, c in best.cols.items() if not f.startswith('_')}, 'output_block_cut': best.cut,
            'alternatives': sorted(alts, key=lambda a: -a['score'])[:5], 'skipped': skipped,
            'blocks': [b.label() for b in taken] if len(taken) > 1 else [],
            'unused': [{'label': b.label(), 'rows': len(b.rows()[0])} for b in unused][:3]}
    for r in rows_out:
        r.pop('_ch_text', None)
    return rows_out, meta


def load_rows(path, sheet=None, text=None):
    """(rows, meta) from a file path, '-' for stdin, or text=."""
    if text is not None or path == '-':
        text = sys.stdin.read() if text is None else text
        if text.count('\t') >= 3 or (text.count(',') >= 6 and re.search(r'(?im)^[^\n]*(input|name|mic)[^\n]*,', text)):
            rows, meta = parse_grid_sheets([('pasted', read_delimited(text))])
            if rows:
                return rows, dict(meta, source='pasted-grid')
        return text_rows(text), {'source': 'text'}
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        rows, meta = parse_grid_sheets(read_xlsx(path), sheet)
        return rows or [], dict(meta, source='xlsx')
    if ext in ('.csv', '.tsv'):
        text = open(path, encoding='utf-8-sig', errors='replace').read()
        rows, meta = parse_grid_sheets([(os.path.basename(path), read_delimited(text))])
        if rows is None:
            return text_rows(text), {'source': 'csv-as-text'}
        return rows, dict(meta, source=ext[1:])
    if ext == '.pdf':
        text = pdf_text(path)
        lines = _pdf_columns(text.replace('\r', '').split('\n'))
        rows = layout_rows(lines)
        if rows:
            return rows, {'source': 'pdf-layout', 'text_chars': len(text)}
        return text_rows(text, pdf=True), {'source': 'pdf-text', 'text_chars': len(text)}
    return load_rows(None, text=open(path, encoding='utf-8', errors='replace').read())


EXTENSIONS = ('.csv', '.tsv', '.txt', '.xlsx', '.xlsm', '.xls', '.docx', '.pdf', '.htm', '.html')


def _text_or_grid(text, source):
    """Pasted or decoded text: a delimited grid when it looks like one, else text lines."""
    skipped, dropped = [], []
    if text.count('\t') >= 3 or (text.count(',') >= 6 and re.search(r'(?im)^[^\n]*(input|name|mic|source)[^\n]*,', text)):
        rows, meta = parse_grid_sheets([(source, read_delimited(text))])
        if rows:
            return rows, dict(meta, source=source + '-grid')
    rows = text_rows(text, skipped=skipped, dropped=dropped)
    return rows, {'source': source, 'skipped': skipped, 'dropped': dropped}


def _best_grid_or_text(sheets, lines, source, sheet=None):
    rows, meta = parse_grid_sheets(sheets, sheet) if sheets else (None, {})
    if rows:
        return rows, dict(meta, source=source)
    if lines:
        return _text_or_grid('\n'.join(lines), source + '-text')
    return [], dict(meta, source=source)


def load_rows_from(filename=None, data=None, text=None, sheet=None):
    """(rows, meta) from an upload (filename + bytes) or pasted text. The extension picks the
    reader; magic bytes override a wrong extension (an .xls that is really HTML, a .txt that is a PDF).
    Raises ValueError with a short message when the file cannot be read."""
    if data is None:
        return _text_or_grid(text or '', 'text')
    ext = os.path.splitext(filename or '')[1].lower()
    head = data[:8]
    if head.startswith(b'%PDF'):
        ext = '.pdf'
    elif head == OLE_MAGIC:
        ext = '.xls' if ext != '.doc' else '.doc'
    elif head.startswith(b'PK'):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            raise ValueError('The file looks like a zip archive but cannot be opened.')
        ext = '.docx' if 'word/document.xml' in names else ('.xlsx' if 'xl/workbook.xml' in names else ext)
    elif ext in ('.xls', '.htm', '.html', '.pdf', '.xlsx', '.xlsm', '.docx') or re.match(rb'\s*<', data[:200]):
        txt = decode_text(data)
        if re.search(r'(?i)<table', txt):
            return _best_grid_or_text(read_html_tables(txt), [], 'html', sheet)
        return _text_or_grid(txt, 'text')
    if ext == '.doc':
        raise ValueError('Old Word .doc files are not supported. Save it as .docx or PDF and try again.')
    if ext in ('.xlsx', '.xlsm'):
        rows, meta = parse_grid_sheets(read_xlsx(io.BytesIO(data)), sheet)
        return rows or [], dict(meta, source='xlsx')
    if ext == '.xls':
        rows, meta = parse_grid_sheets(read_xls(data), sheet)
        return rows or [], dict(meta, source='xls')
    if ext == '.docx':
        tables, lines = read_docx(data)
        return _best_grid_or_text(tables, lines, 'docx', sheet)
    if ext == '.pdf':
        txt = pdf_text(data=data)
        if not txt.strip():
            raise ValueError('The PDF has no text layer (a scan?). Export the list as text, CSV or a spreadsheet.')
        lines = _pdf_columns(txt.replace('\r', '').split('\n'))
        rows = layout_rows(lines)
        if rows:
            return rows, {'source': 'pdf-layout', 'text_chars': len(txt)}
        skipped = []
        return text_rows(txt, pdf=True, skipped=skipped), {'source': 'pdf-text', 'text_chars': len(txt), 'skipped': skipped}
    txt = decode_text(data)
    if ext in ('.csv', '.tsv'):
        rows, meta = parse_grid_sheets([(filename or 'list', read_delimited(txt))])
        if rows:
            return rows, dict(meta, source=ext[1:])
    return _text_or_grid(txt, 'text')


def parse_text(text, **kw):
    rows, meta = load_rows(None, text=text)
    return make_request(rows, meta, **kw)


def parse_file(path, sheet=None, **kw):
    rows, meta = load_rows(path, sheet)
    return make_request(rows, meta, **kw)


def table(req):
    out = []
    for c in req['channels']:
        out.append('%3d %-8s %s %-3s %-16s %-24.24s %-18.18s %.2f %s' % (
            c['number'], c['name'], 'ST' if c['stereo'] else '  ', '48' if c['phantom'] else '',
            c['category'], c['full_name'], c['source']['mic'] or '', c['confidence'], ','.join(c['flags'])))
    return '\n'.join(out)


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('path')
    ap.add_argument('--sheet')
    ap.add_argument('--model')
    ap.add_argument('--firmware')
    ap.add_argument('--rate', type=int)
    ap.add_argument('--artist')
    ap.add_argument('--intent')
    ap.add_argument('--no-stereo', action='store_true')
    ap.add_argument('--naming', choices=('initials', 'fit'), default='initials',
                    help='initials: "Alex AG" -> "A.AG" (initials style); fit: shorten only names over 8')
    ap.add_argument('--table', action='store_true', help='print a readable table instead of JSON')
    a = ap.parse_args(argv)
    kw = dict(stereo=not a.no_stereo, artist=a.artist, model=a.model, firmware=a.firmware, rate=a.rate, intent=a.intent,
              naming=a.naming)
    req = parse_file(a.path, a.sheet, **kw)
    if a.table:
        print(table(req))
        p = req['parse']
        print('# %s' % json.dumps({k: v for k, v in p.items() if k not in ('unparsed', 'alternatives')}))
    else:
        json.dump(req, sys.stdout, indent=1, ensure_ascii=False)
        print()


if __name__ == '__main__':
    main()
